"""
Script to create sample Excel data for testing.
Replace with the real Excel file later.
"""
import openpyxl
import os

def create_sample_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "الطلاب"
    
    # Headers
    ws['A1'] = 'الاسم'
    ws['B1'] = 'الكورس'
    ws['C1'] = 'التاريخ'
    
    # Sample student names (Egyptian/Arabic names with variations to test normalization)
    students = [
        ('أحمد محمد علي',       'دورة تدريبية في البرمجة', '2026-10-15'),
        ('فاطمة عبدالرحمن',     'دورة تدريبية في البرمجة', '2026-10-15'),
        ('محمد إبراهيم حسن',    'دورة تدريبية في البرمجة', '2026-10-15'),
        ('نور الدين عمر',       'دورة تدريبية في البرمجة', '2026-10-15'),
        ('سارة أحمد خالد',      'دورة تدريبية في البرمجة', '2026-10-15'),
        ('عمر عبدالله محمود',   'دورة تدريبية في البرمجة', '2026-10-15'),
        ('إيمان يوسف سعيد',     'دورة تدريبية في البرمجة', '2026-10-15'),
        ('علي حسين الشريف',     'دورة تدريبية في البرمجة', '2026-10-15'),
        ('مريم إسماعيل فاروق',  'دورة تدريبية في البرمجة', '2026-10-15'),
        ('خالد عبدالعزيز نور',  'دورة تدريبية في البرمجة', '2026-10-15'),
        ('هبة الله محمد رشاد',  'دورة تدريبية في البرمجة', '2026-10-15'),
        ('عبدالرحمن طارق فتحي', 'دورة تدريبية في البرمجة', '2026-10-15'),
    ]
    
    for row_data in students:
        ws.append(row_data)
    
    os.makedirs('data', exist_ok=True)
    wb.save('data/students.xlsx')
    print("Done: data/students.xlsx created successfully.")

if __name__ == '__main__':
    create_sample_excel()
